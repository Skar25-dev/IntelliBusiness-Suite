from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.config import settings
from app.services.report_service import ReportService, get_columns_for_model


def _rows(path):
    sheet = load_workbook(path).active
    values = list(sheet.iter_rows(values_only=True))
    return list(values[0]), [list(r) for r in values[1:]]


def test_get_columns_for_model_excludes_id():
    columns = get_columns_for_model("ventas")
    assert "quantity" in columns and "id" not in columns
    assert "current_stock" in get_columns_for_model("inventario")


def test_unknown_report_type_raises(db_session):
    with pytest.raises(ValueError):
        ReportService(db_session).get_report(report_type="no_existe")


def test_sales_report_respects_30_day_window(db_session):
    path = ReportService(db_session).get_report(report_type="ventas")

    # Solo "product_name" y "date" y "amount" existen en Sale ("category" se descarta)
    headers, rows = _rows(path)
    assert headers == ["PRODUCT NAME", "AMOUNT", "DATE"]
    # 2 productos x 29 días (edades 1..29 caen dentro de 30 días)
    assert len(rows) == 2 * 29


def test_days_parameter_widens_the_window(db_session):
    path = ReportService(db_session).get_report(report_type="ventas", days=60)
    _, rows = _rows(path)
    assert len(rows) == 2 * 40


def test_sales_report_category_filter_uses_product_category(db_session):
    """Sale no tiene columna `category`: se filtra por la categoría de su producto."""
    path = ReportService(db_session).get_report(report_type="ventas", category="Hogar", days=60)
    _, rows = _rows(path)
    assert len(rows) == 40
    assert {r[0] for r in rows} == {"Silla"}


def test_inventory_report_and_category_filter(db_session, monkeypatch):
    monkeypatch.setattr(
        settings, "report_settings",
        {"default_days": 30, "included_columns": ["name", "category", "current_stock"],
         "sheet_name": "Inventario"},
    )
    service = ReportService(db_session)

    _, all_rows = _rows(service.get_report(report_type="inventario"))
    _, home_rows = _rows(service.get_report(report_type="inventario", category="Hogar"))

    assert len(all_rows) == 2
    assert home_rows == [["Silla", "Hogar", 5]]


def test_no_data_returns_none(db_session):
    assert ReportService(db_session).get_report(report_type="inventario", category="Inexistente") is None


def test_pdf_report_is_generated(db_session, monkeypatch):
    monkeypatch.setattr(settings, "report_format", "pdf")
    path = ReportService(db_session).get_report(report_type="ventas")

    assert str(path).endswith(".pdf")
    assert Path(path).read_bytes().startswith(b"%PDF")
